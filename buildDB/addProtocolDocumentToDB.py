import os
import sqlite3
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

def iter_block_items(parent):
    """
    Yield each paragraph and table child within *parent*, in document order.
    """
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Parent must be a Document or _Cell object")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def populate_protocols_from_docx(db_path, docs_directory):
    """
    Reads all .docx files from a directory, extracts their content (converting
    tables to bullet points), and inserts the data into the 'protocols' table
    of a SQLite database.

    Args:
        db_path (str): The path to the SQLite database file.
        docs_directory (str): The path to the directory containing the .docx files.

    Returns:
        bool: True if successful, False otherwise.
    """
    if not os.path.isdir(docs_directory):
        print(f"Error: Directory not found at '{docs_directory}'")
        return False

    try:
        # Connect to the SQLite database
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        print(f"Successfully connected to database: {db_path}")

        # Create the 'protocols' table if it doesn't exist
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS protocols (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                experiment TEXT NOT NULL UNIQUE,
                content TEXT
            )
        ''')
        print("Table 'protocols' is ready.")

        # Get a list of all .docx files
        filenames = sorted([f for f in os.listdir(docs_directory) if f.lower().endswith('.docx')])
        
        if not filenames:
            print(f"No .docx files found in '{docs_directory}'.")
            return True

        print(f"Found {len(filenames)} .docx files to process.")

        for filename in filenames:
            file_path = os.path.join(docs_directory, filename)
            print(f"  - Processing '{filename}'...")
            
            try:
                source_doc = Document(file_path)
                full_content_parts = []
                
                # Extract content from paragraphs and tables
                for block in iter_block_items(source_doc):
                    if isinstance(block, Paragraph):
                        full_content_parts.append(block.text)
                    elif isinstance(block, Table):
                        # Add a marker for the start of a table
                        full_content_parts.append("\n--- TABLE DATA ---")
                        for row in block.rows:
                            row_text = "\t".join(cell.text for cell in row.cells)
                            # Format as a bullet point
                            full_content_parts.append(f"  • {row_text}")
                        full_content_parts.append("--- END TABLE ---\n")

                # Join all parts into a single string
                final_content = "\n".join(full_content_parts)
                
                experiment_name, _ = os.path.splitext(filename)
                # Insert or Replace the data into the table.
                # 'OR REPLACE' ensures that if you run the script again, it updates the entry
                # for a file instead of creating a duplicate or erroring.
                cursor.execute(
                    "INSERT OR REPLACE INTO protocols (experiment, content) VALUES (?, ?)",
                    (experiment_name, final_content)
                )
                print(f"    - Added/Updated content for '{filename}' in the database.")

            except Exception as e:
                print(f"    Warning: Could not process file '{filename}'. Error: {e}. Skipping.")

        # Commit the changes and close the connection
        conn.commit()
        print("\nAll changes have been committed to the database.")
        return True

    except sqlite3.Error as e:
        print(f"Database error: {e}")
        return False
    finally:
        if conn:
            conn.close()
            print("Database connection closed.")

# --- Main Execution ---
if __name__ == "__main__":
    # To run this script, you first need to install the required library:
    # pip install python-docx

    # --- Configuration ---
    database_file = 'experiments.db'
    source_directory = '/Users/zyu7/Documents/KRgpt/data/SOPs/ChangeDB'

    # --- Create a dummy directory and files for a runnable example ---
    if not os.path.exists(source_directory):
        print(f"Creating dummy directory '{source_directory}' for demonstration.")
        os.makedirs(source_directory)
        
        # Create dummy file A with a table
        doc_a = Document()
        doc_a.add_paragraph("This is text from file_A.docx, before the table.")
        table = doc_a.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Header 1'
        table.cell(0, 1).text = 'Header 2'
        table.cell(1, 0).text = 'Cell A'
        table.cell(1, 1).text = 'Cell B'
        doc_a.add_paragraph("This is text after the table.")
        doc_a.save(os.path.join(source_directory, "file_A_with_table.docx"))

        # Create dummy file B
        doc_b = Document()
        doc_b.add_paragraph("This is the content from the second file, file_B.docx.")
        doc_b.save(os.path.join(source_directory, "file_B.docx"))
        print("Dummy files created.")
    # --- End of dummy file setup ---

    # Run the main function
    success = populate_protocols_from_docx(database_file, source_directory)

    if success:
        print("\nScript finished successfully!")
    else:
        print("\nScript finished with errors.")
