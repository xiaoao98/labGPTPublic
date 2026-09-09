import os
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

def iter_block_items(parent):
    """
    Yield each paragraph and table child within *parent*, in document order.
    Each returned value is an instance of either Paragraph or Table.
    """
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Parent must be a Document or _Cell object")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def concatenate_files_in_directory(directory_path, output_filepath):
    """
    Reads all .docx files in a given directory and concatenates their content,
    including paragraphs and tables, into a single Word (.docx) file.
    Table content is converted to bulleted text.

    Args:
        directory_path (str): The path to the directory containing the .docx files.
        output_filepath (str): The path for the final concatenated Word document.

    Returns:
        bool: True if successful, False otherwise.
    """
    if not os.path.isdir(directory_path):
        print(f"Error: Directory not found at '{directory_path}'")
        return False

    try:
        # Create a new Word document object for the output
        main_document = Document()
        main_document.add_heading('Concatenated Document', level=0)

        # Get a list of all files ending with .docx in the directory.
        filenames = sorted([f for f in os.listdir(directory_path) if f.lower().endswith('.docx')])
        
        if not filenames:
            print(f"No .docx files found in '{directory_path}'. An empty Word document will be created.")
            main_document.save(output_filepath)
            return True

        print(f"Found {len(filenames)} files to concatenate.")

        for i, filename in enumerate(filenames):
            file_path = os.path.join(directory_path, filename)
            
            print(f"  - Processing '{filename}'...")
            
            try:
                source_doc = Document(file_path)
                
                # Use the helper function to iterate over paragraphs and tables
                for block in iter_block_items(source_doc):
                    if isinstance(block, Paragraph):
                        # Copy paragraph text
                        main_document.add_paragraph(block.text)
                    elif isinstance(block, Table):
                        # Convert table content to bullet points
                        print(f"    - Converting a table to text...")
                        for row in block.rows:
                            # Join the text of all cells in a row, separated by tabs
                            row_text = "\t".join(cell.text for cell in row.cells)
                            # Add the row's content as a bullet point
                            main_document.add_paragraph(row_text, style='List Bullet')
                
                # Add a simple newline for a little space between documents
                if i < len(filenames) - 1:
                    main_document.add_paragraph("\n")

            except Exception as e:
                print(f"    Warning: Could not process file '{filename}'. Error: {e}. Skipping.")

        # Save the final Word document
        main_document.save(output_filepath)
        print(f"\nSuccessfully concatenated all files into '{output_filepath}'.")
        return True

    except Exception as e:
        print(f"An unexpected error occurred: {e}")
        return False

# --- Main Execution ---
if __name__ == "__main__":
    # To run this script, you first need to install the required library:
    # pip install python-docx

    # --- Configuration ---
    source_directory = '/Users/zyu7/Documents/KRgpt/data/SOPs/categorized/Genotyping'
    concatenated_output_file = 'concatenated_document.docx'

    # --- Create a dummy directory and files for a runnable example ---
    if not os.path.exists(source_directory):
        print(f"Creating dummy directory '{source_directory}' for demonstration.")
        os.makedirs(source_directory)
        
        # Create dummy file A with a table
        doc_a = Document()
        doc_a.add_paragraph("This is the content of file_A.docx, before the table.")
        table = doc_a.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Header 1'
        table.cell(0, 1).text = 'Header 2'
        table.cell(1, 0).text = 'Cell A'
        table.cell(1, 1).text = 'Cell B'
        doc_a.add_paragraph("This is the content after the table.")
        doc_a.save(os.path.join(source_directory, "file_A_with_table.docx"))

        # Create dummy file B
        doc_b = Document()
        doc_b.add_paragraph("This is the content from the second file, file_B.docx.")
        doc_b.save(os.path.join(source_directory, "file_B.docx"))
        print("Dummy files created.")
    # --- End of dummy file setup ---

    # Run the main function
    success = concatenate_files_in_directory(source_directory, concatenated_output_file)

    if success:
        print(f"\nScript finished successfully! Check '{concatenated_output_file}' for the result.")
    else:
        print("\nScript finished with errors.")
